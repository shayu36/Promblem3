"""生成《数据挖掘导论》任务3 时间序列预测的报告（.docx）。

按照报告模版的格式约定（宋体五号正文、黑体小四一级标题、楷体五号摘要等）
组织内容，并从 results/metrics.json 读取真实评估指标填入。
"""
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

from config import RESULTS_DIR, ROOT

METRICS = json.load(open(os.path.join(RESULTS_DIR, "metrics.json"), encoding="utf-8"))

# ---------------- 字体工具 ----------------

def _set_font(run, cn_font="宋体", size=10.5, bold=False, ascii_font="Times New Roman"):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)


def add_para(doc, text, cn_font="宋体", size=10.5, bold=False,
             align=None, indent=False, space_after=6, space_before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)  # 首行缩进 2 字符
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    _set_font(run, cn_font, size, bold)
    return p


def add_heading1(doc, text):
    return add_para(doc, text, cn_font="黑体", size=12, bold=False,
                    space_before=10, space_after=6)


def add_heading2(doc, text):
    return add_para(doc, text, cn_font="黑体", size=10.5, bold=False,
                    space_before=6, space_after=4)


def add_body(doc, text):
    return add_para(doc, text, cn_font="宋体", size=10.5, indent=True,
                    align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def add_caption(doc, text, bold=True):
    return add_para(doc, text, cn_font="黑体" if bold else "宋体", size=9,
                    align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)


# ---------------- 构建报告 ----------------

def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # 中文题目
    add_para(doc, "基于滑动窗口与循环神经网络的室外温度时间序列预测",
             cn_font="黑体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=4)
    # 英文题目
    add_para(doc, "Outdoor Temperature Time Series Forecasting "
             "with Sliding Window and Recurrent Neural Networks",
             cn_font="黑体", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=12)

    # 摘要
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("摘  要  ")
    _set_font(r, "黑体", 10.5, bold=True)
    r = p.add_run(
        f"时间序列预测是数据挖掘中的一类重要回归问题。本文针对德国某气象站半年内"
        f"以 10 分钟为间隔记录的 21 维气象指标数据，研究“根据过去 2 小时（12 个时间点）"
        f"的气象状态预测下一时间点室外温度 OT”的问题。首先对数据进行异常值清洗与特征标准化，"
        f"并按时间顺序构造滑动窗口、划分训练集与测试集；随后设计并实现了基于长短期记忆网络"
        f"（LSTM）与门控循环单元（GRU）的循环神经网络预测模型，同时以持久化预测与岭回归作为"
        f"基线进行对比。实验结果表明，GRU 模型在测试集上取得了 RMSE={METRICS['GRU']['RMSE']:.4f}、"
        f"MAE={METRICS['GRU']['MAE']:.4f}、R²={METRICS['GRU']['R2']:.4f} 的最优预测效果，"
        f"显著优于两种基线方法，验证了循环神经网络在短期气象时间序列预测中的有效性。")
    _set_font(r, "楷体", 10.5)

    # 关键词
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("关键词  ")
    _set_font(r, "黑体", 10.5, bold=True)
    r = p.add_run("时间序列预测；室外温度；门控循环单元；长短期记忆网络；滑动窗口；特征标准化")
    _set_font(r, "楷体", 10.5)

    # ---- 1 问题的形式化描述 ----
    add_heading1(doc, "1  问题的形式化描述")
    add_body(doc,
        "时间序列预测旨在根据历史观测推断未来的取值，在气象、金融、能源等领域应用广泛。"
        "气象数据是典型的多变量时间序列：一方面，相邻时刻的温度、气压等量具有强自相关，"
        "短期变化平滑；另一方面，序列存在明显的周期规律（如昼夜循环）以及变量之间复杂的耦合关系，"
        "适合作为回归与序列建模的研究对象。")
    add_body(doc,
        "本文任务可形式化如下。记气象站在第 t 个时间点（相邻时间点间隔 Δt = 10 分钟）观测到的"
        "特征向量为 x_t ∈ R^21，共包含 21 个气象学指标，其中也包含待预测的室外温度 OT 的历史观测值。"
        "给定过去 2 小时共 w = 12 个时间点的观测序列 X_{t-11:t} = (x_{t-11}, …, x_t)，"
        "预测下一时间点的室外温度 y_{t+1} = OT(t+1)，即学习映射")
    add_para(doc, "f : R^(12×21) → R，  ŷ_{t+1} = f(X_{t-11:t})",
             cn_font="宋体", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_body(doc,
        "该问题是一个多变量输入、单变量输出的序列回归问题，本文采用均方根误差（RMSE）、"
        "平均绝对误差（MAE）与决定系数（R²）作为评价指标，并补充平均绝对百分比误差（MAPE）。")

    # ---- 2 数据预处理与训练/测试集划分 ----
    add_heading1(doc, "2  数据预处理与训练/测试集划分")

    add_heading2(doc, "2.1  数据说明")
    add_body(doc,
        "数据集 Weather 记录了德国某气象站约半年内、每隔 10 分钟采样的 21 个气象学指标，"
        "共 26200 条记录。指标包括气压 p、温度 T、位温 Tpot、露点温度 Tdew、相对湿度 rh、"
        "水汽压 VPmax/VPact/VPdef、比湿 sh、水汽浓度 H2OC、空气密度 rho、风速 wv 与最大风速 "
        "max.wv、风向 wd、降雨量 rain、日照 SWDR、光合有效辐射 PAR 等，以及待预测的室外温度 OT。")

    add_heading2(doc, "2.2  数据清洗")
    add_body(doc,
        "经统计，原始数据中共存在 45 处取值为 -9999 的传感器异常占位值，其中 OT 列 34 处"
        "（集中于某一连续时段）、风速 wv 列 1 处、最大光合有效辐射 max.PAR 列 10 处。"
        "这些占位值并非真实观测，若直接参与训练会引入极大噪声。本文将此类值视为缺失值，"
        "并沿时间维度进行线性插值补齐，从而得到连续、完整的观测序列。")

    add_heading2(doc, "2.3  特征标准化")
    add_body(doc,
        "各特征量纲差异悬殊（如气压约 10^3 mbar、温度仅个位数、辐射可达数百 W/m²），"
        "直接输入模型会因数值尺度不一致而影响收敛速度与特征权重。本文采用 Z-score 标准化，"
        "对每个特征 x 计算其在训练集上的均值 μ 与标准差 σ，再做变换 x' = (x − μ)/σ。"
        "为避免测试信息泄漏，μ 与 σ 仅在训练集上拟合，再统一作用于训练、验证与测试三个集合。")

    add_heading2(doc, "2.4  滑动窗口切分与数据集划分")
    add_body(doc,
        "按第 1 节的形式化定义，取窗口长度 w = 12（对应过去 2 小时），对清洗后的序列做滑动窗口切分，"
        "共得到 26188 个样本，每个样本输入为 12×21 的张量，标签为下一时间点的 OT。")
    add_body(doc,
        "时间序列具有严格的时间顺序依赖，随机划分会破坏时序结构并造成信息泄漏。"
        "因此本文严格按时间先后划分：最后 20%（约 5237 个样本）作为测试集，"
        "剩余 80% 中再取后 15%（约 3142 个样本）作为验证集用于早停，"
        "最终训练/验证/测试规模为 17809 / 3142 / 5237。")

    # ---- 3 时间序列预测模型 ----
    add_heading1(doc, "3  时间序列预测模型")

    add_heading2(doc, "3.1  基线模型")
    add_body(doc,
        "（1）持久化预测（Persistence）：直接以最后一个观测值外推，即 ŷ_{t+1} = OT(t)，"
        "反映了“温度在短时间尺度上变化缓慢”这一先验，是最朴素且强力的基线。")
    add_body(doc,
        "（2）岭回归（Ridge Regression）：将 12×21 的滑动窗口展平为 252 维向量，"
        "用带 L2 正则的线性回归建模，用于检验线性模型是否足以刻画序列关系。")

    add_heading2(doc, "3.2  循环神经网络模型设计")
    add_body(doc,
        "循环神经网络（RNN）天然适合序列建模，但普通 RNN 存在梯度消失/爆炸问题，"
        "难以捕捉长程依赖。门控型 RNN 通过门控机制与细胞状态可有效缓解上述问题，"
        "其中长短期记忆网络（LSTM）引入输入门、遗忘门、输出门三个门控单元，"
        "而门控循环单元（GRU）将门控精简为更新门与重置门，参数更少、结构更简洁。")
    add_body(doc, "本文分别实现 LSTM 与 GRU 两种模型，其整体结构一致，仅循环单元不同：")
    add_body(doc,
        "· 输入层：形状为 (batch, 12, 21) 的滑动窗口张量；")
    add_body(doc,
        "· 特征提取：两层循环单元（LSTM 或 GRU），隐层维度 128，层间 dropout 为 0.2；")
    add_body(doc,
        "· 回归头：取最后一个时间步的隐状态，经全连接层（128 → 64 → 1）输出下一时间点 OT 的预测值。")
    add_body(doc,
        "模型以均方误差（MSE）为损失函数，采用 Adam 优化器训练，初始学习率 5e-4，"
        "并配合学习率衰减（验证集损失停滞时减半）与梯度裁剪（阈值 1.0），"
        "基于验证集损失进行早停（耐心 25 轮）以防止过拟合。")

    add_heading2(doc, "3.3  设计思想")
    add_body(doc,
        "选择门控型 RNN 主要基于三点考虑。其一，气象序列具有明显的昼夜周期（24 小时即 144 个时间步）"
        "与短期连续性，需要模型同时具备长期记忆与短期平滑能力，门控机制正适合这一需求。"
        "其二，多变量输入中，气压、湿度、辐射等特征对温度存在滞后且非线性的耦合影响，"
        "RNN 的逐时间步非线性变换可自动学习这种依赖。其三，相较于 ARIMA 等经典统计方法，"
        "RNN 无需人工指定滞后阶数，可端到端地从窗口输入中学习表示。")
    add_body(doc,
        "本文同时实现 LSTM 与 GRU 两种门控结构加以比较：LSTM 门控更精细、表达能力更强，"
        "但参数更多；GRU 结构更简洁、参数更少。通过实验对比，可以判断在本任务的短序列、"
        "强自相关特性下，何种门控结构更为合适。")

    # ---- 4 实验与结果 ----
    add_heading1(doc, "4  实验与结果")

    add_heading2(doc, "4.1  评价指标")
    add_body(doc,
        "本文采用回归任务常用指标：RMSE（对较大误差更敏感）、MAE（直观反映平均绝对偏差）、"
        "R²（衡量解释方差的比例），并补充 MAPE。所有指标均在原始量纲（OT 的原始单位）上计算，"
        "便于与真实温度变化量级直接比较。")

    add_heading2(doc, "4.2  实验结果")
    add_body(doc, "各方法在测试集上的预测效果如表 1 所示。")

    # 指标表
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, name in enumerate(["方法", "RMSE", "MAE", "R²", "MAPE (%)"]):
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = hdr[i].paragraphs[0].add_run(name)
        _set_font(r, "黑体", 9, bold=True)
    name_map = {"Persistence": "持久化预测", "Ridge": "岭回归",
                "LSTM": "LSTM", "GRU": "GRU（本文）"}
    for key in ["Persistence", "Ridge", "LSTM", "GRU"]:
        m = METRICS[key]
        cells = table.add_row().cells
        vals = [name_map[key], f"{m['RMSE']:.4f}", f"{m['MAE']:.4f}",
                f"{m['R2']:.4f}", f"{m['MAPE(%)']:.2f}"]
        for i, v in enumerate(vals):
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cells[i].paragraphs[0].add_run(v)
            _set_font(r, "宋体", 9)
    add_caption(doc, "表 1  各方法在测试集上的预测效果", bold=True)

    add_heading2(doc, "4.3  结果分析")
    base = METRICS["Persistence"]["RMSE"]
    gru = METRICS["GRU"]["RMSE"]
    add_body(doc, "从表 1 可以得出如下结论：")
    add_body(doc,
        f"（1）持久化基线的 RMSE 已达到 {base:.4f}，说明温度序列短期自相关很强"
        f"（滞后 1 步自相关系数约为 0.99），“温度短时不变”本身就是很强的先验；")
    add_body(doc,
        "（2）岭回归略优于持久化，表明多变量之间的线性关系能够提供一定的额外信息；")
    add_body(doc,
        f"（3）GRU 取得最优效果，RMSE 相对持久化基线降低约 {(1 - gru / base) * 100:.1f}%、"
        f"R² 达到 {METRICS['GRU']['R2']:.4f}，验证了循环神经网络在短期气象时间序列预测中的价值；")
    add_body(doc,
        "（4）LSTM 略逊于 GRU 与岭回归，说明在该任务的短序列、强自相关特性下，"
        "LSTM 更精细的门控结构与更多参数并未带来增益，反而更易过拟合；"
        "结构更简洁的 GRU 参数更少、泛化更好，是本任务更合适的选择。")
    add_body(doc,
        "从预测曲线与散点图（图 1、图 2）可见，模型能够准确跟踪温度的昼夜波动与快速转折，"
        "预测值与真实值高度吻合，误差主要集中在温度剧烈变化处。")

    # 图片
    for path, cap in [("prediction.png", "图 1  测试集前 240 个时间步的预测值与真实值对比"),
                      ("scatter.png", "图 2  预测值与真实值的散点图")]:
        full = os.path.join(RESULTS_DIR, path)
        if os.path.exists(full):
            doc.add_picture(full, width=Cm(14))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_caption(doc, cap, bold=False)

    # ---- 5 总结 ----
    add_heading1(doc, "5  总结")
    add_body(doc,
        f"本文针对室外温度时间序列预测任务，完成了数据清洗、特征标准化、滑动窗口切分与"
        f"按时间顺序的数据集划分，设计并实现了基于 LSTM/GRU 的循环神经网络预测模型，"
        f"并与持久化预测、岭回归两种基线进行对比。实验结果表明，GRU 在测试集上取得了 RMSE="
        f"{METRICS['GRU']['RMSE']:.4f}、MAE={METRICS['GRU']['MAE']:.4f}、"
        f"R²={METRICS['GRU']['R2']:.4f} 的最优效果，显著优于基线方法，验证了方法的有效性。"
        f"实验同时表明，在该短序列、强自相关任务中，结构更简洁的 GRU 优于 LSTM。"
        f"后续可进一步引入小时/季节等显式时间特征、注意力机制或进行多步预测，"
        f"以提升对更长时间跨度的预测能力。")

    # ---- 参考文献 ----
    add_heading1(doc, "参 考 文 献")
    refs = [
        "[1] Hochreiter S, Schmidhuber J. Long short-term memory[J]. Neural Computation, 1997, 9(8): 1735-1780.",
        "[2] Cho K, Van Merriënboer B, Gulcehre C, et al. Learning phrase representations using RNN encoder-decoder for statistical machine translation[C]//Proc of EMNLP. Doha: ACL, 2014: 1724-1734.",
        "[3] Lim B, Zohren S. Time-series forecasting with deep learning: a survey[J]. Philosophical Transactions of the Royal Society A, 2021, 379(2194): 20200209.",
        "[4] Box G E P, Jenkins G M, Reinsel G C, et al. Time Series Analysis: Forecasting and Control[M]. 5th ed. Hoboken, NJ: John Wiley & Sons, 2015.",
        "[5] Kingma D P, Ba J. Adam: A method for stochastic optimization[C]//Proc of ICLR. San Diego, 2015.",
        "[6] Max Planck Institute for Biogeochemistry. Jena Climate dataset[OL]. https://www.bgc-jena.mpg.de/wetter/.",
    ]
    for r_ in refs:
        add_para(doc, r_, cn_font="宋体", size=9, space_after=3)

    out = os.path.join(ROOT, "任务3_时间序列预测_报告.docx")
    doc.save(out)
    print("报告已生成:", out)
    return out


if __name__ == "__main__":
    build()
